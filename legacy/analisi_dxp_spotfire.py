import csv
import gzip
import hashlib
import re
import shutil
import sys
import zipfile
import zlib
import xml.etree.ElementTree as ET
from collections import Counter, OrderedDict
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None


CARTELLA_BASE = Path("input")

FILE_DXP = []

CARTELLA_OUTPUT = Path("output")

ESTRAI_ARCHIVIO_COMPLETO = True

DIMENSIONE_MASSIMA_TESTO = 50 * 1024 * 1024

ESTENSIONI_BINARIE = {".sbdf", ".stdf", ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".ico", ".dll", ".exe", ".zip", ".pdf"}

ESTENSIONI_IMMAGINE = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".ico", ".svg"}

TIPI_ISTANTANEA = re.compile(r"Bookmark|Snapshot|Undo|History")

CHIAVI_CODICE = {"scriptcode", "code", "script", "scripttext", "source", "sourcecode", "body", "text", "content"}

PROPRIETA_SCHEDA_GENERAL = {"description", "keywords", "allowwebplayerresume", "publicbookmarkcreation", "privatebookmarkcreation", "name", "title", "author", "comments"}

CHIAVI_NOME = ("name", "title", "displayname", "scriptname", "caption", "titleexpression")


TIPI_VISUAL = {
    "BarChart", "LineChart", "CombinationChart", "CrossTablePlot", "TablePlot", "SummaryTable", "ScatterPlot",
    "ScatterPlot3D", "PieChart", "Treemap", "TreeMap", "HeatMap", "BoxPlot", "GraphicalTable", "KpiChart",
    "KPIChart", "WaterfallChart", "ParallelCoordinatePlot", "MapChart", "HtmlTextArea", "TextArea",
}

TIPI_TEXT_AREA = {"HtmlTextArea", "TextArea"}

SEGNALI_PYTHON = [re.compile(x, re.M) for x in (
    r"^\s*from\s+\w[\w.]*\s+import\b",
    r"^\s*import\s+\w",
    r"\bDocument\.\w",
    r"\bApplication\.\w",
    r"\bclr\.",
    r"^\s*def\s+\w+\s*\(",
    r"^\s*for\s+\w+\s+in\s+.+:\s*$",
    r"^\s*if\s+.+:\s*$",
    r"\.As\[\w+\]",
    r"\bSpotfire\.Dxp\.\w",
)]

SEGNALI_JS = [re.compile(x, re.M) for x in (
    r"\bfunction\s*\w*\s*\(",
    r"\bvar\s+\w",
    r"\b(?:let|const)\s+\w",
    r"\bdocument\.\w",
    r"\$\(",
    r"\bjQuery\b",
    r"\bset(?:Timeout|Interval)\s*\(",
    r"addEventListener",
    r"=>",
    r"getElementById|querySelector",
    r"\bconsole\.\w",
    r"\bwindow\.\w",
)]




def impronta(testo):
    return hashlib.md5(re.sub(r"\s+", "", testo or "").encode("utf-8", "ignore")).hexdigest()


def pulisci_nome(testo, lunghezza=80):
    pulito = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "_", str(testo or "")).strip().strip(".")
    pulito = re.sub(r"\s+", " ", pulito)
    return pulito[:lunghezza].strip() or "senza nome"


def tipo_base(tipo):
    return (tipo or "").split(",")[0].strip()


def ultimo_segmento(tipo):
    base = tipo_base(tipo).split("`")[0]
    return re.split(r"[.+]", base)[-1] if base else ""


def chiave_campo(nome):
    return re.sub(r"^(?:m_|_)+", "", (nome or "").lower())


def normalizza_percorso(percorso):
    return (percorso or "").replace("\\", "/").lstrip("/").lower()


def testo_leggibile(testo):
    campione = testo[:4000]
    if not campione:
        return False
    buoni = sum(1 for c in campione if (ord(c) < 0x250 and (c.isprintable() or c in "\r\n\t")))
    return buoni / len(campione) > 0.9


def decodifica_bytes(raw):
    dati = raw
    if dati[:2] == b"\x1f\x8b":
        try:
            dati = gzip.decompress(dati)
        except Exception:
            pass
    elif dati[:1] == b"\x78":
        try:
            dati = zlib.decompress(dati)
        except Exception:
            pass
    if dati[:3] == b"\xef\xbb\xbf":
        codifiche = ["utf-8-sig"]
    elif dati[:2] in (b"\xff\xfe", b"\xfe\xff"):
        codifiche = ["utf-16"]
    else:
        codifiche = ["utf-8", "utf-16", "cp1252"]
    for codifica in codifiche:
        try:
            testo = dati.decode(codifica)
        except (UnicodeDecodeError, LookupError):
            continue
        if testo_leggibile(testo):
            return testo
    return None


def punti_codice(testo):
    campione = (testo or "")[:200000]
    return (
        sum(1 for r in SEGNALI_PYTHON if r.search(campione)),
        sum(1 for r in SEGNALI_JS if r.search(campione)),
    )


def indovina_linguaggio(testo, rigoroso=True):
    if not testo or not testo.strip():
        return None
    punti_py, punti_js = punti_codice(testo)
    if not rigoroso:
        return "JavaScript" if punti_js > punti_py else "IronPython"
    if punti_py >= 3 and punti_py > punti_js:
        return "IronPython"
    if punti_js >= 3 and punti_js > punti_py and ("{" in testo or ";" in testo):
        return "JavaScript"
    return None


def linguaggio_dichiarato(dichiarato):
    valore = (dichiarato or "").strip().lower()
    if "javascript" in valore or valore in {"js", "jscript"}:
        return "JavaScript"
    if "ironpython" in valore:
        return "IronPython"
    if "python" in valore:
        return "Python"
    return None



def scrivi_testo(percorso, contenuto, errori=None):
    try:
        percorso.parent.mkdir(parents=True, exist_ok=True)
        with open(percorso, "w", encoding="utf-8", newline="") as f:
            f.write(contenuto or "")
        return True
    except OSError as e:
        if errori is not None:
            errori.append(f"Impossibile scrivere {percorso}: {e}")
        return False



def testo_cella(valore):
    if valore is None:
        return ""
    if isinstance(valore, (int, float)):
        return valore
    testo = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(valore))
    return testo[:32000]


def scrivi_tabella(percorso, intestazione, righe):
    percorso.parent.mkdir(parents=True, exist_ok=True)
    righe = list(righe)
    if Workbook is None:
        with open(percorso.with_suffix(".csv"), "w", encoding="utf-8-sig", newline="") as f:
            scrittore = csv.writer(f, delimiter=";")
            scrittore.writerow(intestazione)
            for riga in righe:
                scrittore.writerow(riga)
        return
    cartella_lavoro = Workbook()
    foglio = cartella_lavoro.active
    foglio.title = pulisci_nome(percorso.stem, 31).replace("[", "").replace("]", "")
    foglio.append(list(intestazione))
    for cella in foglio[1]:
        cella.font = Font(bold=True)
    for riga in righe:
        foglio.append([testo_cella(v) for v in riga])
    foglio.freeze_panes = "A2"
    if righe:
        foglio.auto_filter.ref = foglio.dimensions
    for indice_colonna, titolo in enumerate(intestazione, 1):
        lunghezze = [len(str(titolo))] + [len(str(r[indice_colonna - 1])) for r in righe[:200] if indice_colonna - 1 < len(r)]
        foglio.column_dimensions[get_column_letter(indice_colonna)].width = min(60, max(10, max(lunghezze) + 2))
    cartella_lavoro.save(percorso.with_suffix(".xlsx"))


def nome_file_libero(nome, estensione, usati):
    base = pulisci_nome(nome)
    candidato = base
    progressivo = 2
    while (candidato + estensione).lower() in usati:
        candidato = f"{base} ({progressivo})"
        progressivo += 1
    usati.add((candidato + estensione).lower())
    return candidato + estensione



def e_visual(tipo):
    base = tipo_base(tipo)
    return "+" not in base and base.startswith("Spotfire.Dxp.Application.Visuals") and ultimo_segmento(tipo) in TIPI_VISUAL


class AnalizzatoreDxp:

    def __init__(self, percorso):
        self.percorso = Path(percorso)
        self.nome = self.percorso.stem
        self.xml = OrderedDict()
        self.indici = OrderedDict()
        self.testi = {}
        self.testi_normalizzati = {}
        self.voci_zip = []
        self.immagini = []
        self.risorse = []
        self.risorse_per_id = {}
        self.risorse_per_nome = {}
        self.riferimenti_risorse = {}
        self.cache_tipo = {}
        self.cache_istantanea = {}
        self.tipi_vivi = Counter()
        self.tipi_istantanea = Counter()
        self.oggetti_script = []
        self.oggetti_data_function = []
        self.oggetti_proprieta = []
        self.oggetti_colonne_calcolate = []
        self.segnalibri = 0
        self.pagine = []
        self.visual = Counter()
        self.text_area = []
        self.script = []
        self.data_function = []
        self.hash_data_function = set()
        self.proprieta = []
        self.proprieta_documento = []
        self.proprieta_documento_standard = []
        self.proprieta_citate = {}
        self.colonne_calcolate = []
        self.query = []
        self.nomi_query = set()
        self.hash_query = set()
        self.connessioni = []
        self.connessioni_viste = set()
        self.copie_documento = []
        self.file_principale = None
        self.nome_risorsa_per_file = {}
        self.tabelle_sorgente = []
        self.tabelle_viste = set()
        self.avvisi = []
        self.errori = []

    def esegui(self):
        self.carica_archivio()
        self.indicizza()
        self.censisci_oggetti()
        self.estrai_pagine_e_visual()
        self.estrai_data_function()
        self.estrai_script()
        self.estrai_proprieta()
        self.estrai_colonne_calcolate()
        self.estrai_query_e_connessioni()
        self.estrai_tabelle_sorgente()

    def parse_xml(self, dati):
        try:
            if isinstance(dati, bytes):
                radice = ET.fromstring(dati)
            else:
                radice = ET.fromstring(re.sub(r"^\s*<\?xml[^>]*\?>", "", dati))
        except (ET.ParseError, ValueError):
            if isinstance(dati, bytes):
                testo = decodifica_bytes(dati)
                if testo:
                    return self.parse_xml(testo)
            return None
        for elemento in radice.iter():
            if isinstance(elemento.tag, str) and "}" in elemento.tag:
                elemento.tag = elemento.tag.rsplit("}", 1)[1]
            if any("}" in k for k in elemento.attrib):
                elemento.attrib = {k.rsplit("}", 1)[-1]: v for k, v in elemento.attrib.items()}
        return radice

    def carica_archivio(self):
        with zipfile.ZipFile(self.percorso, "r") as zf:
            for info in zf.infolist():
                if info.is_dir():
                    continue
                nome = info.filename
                self.voci_zip.append((nome, info.file_size, info.compress_size))
                estensione = Path(nome).suffix.lower()
                if estensione in ESTENSIONI_IMMAGINE:
                    self.immagini.append(nome)
                if estensione in ESTENSIONI_BINARIE or info.file_size > DIMENSIONE_MASSIMA_TESTO:
                    continue
                try:
                    raw = zf.read(info)
                except Exception as e:
                    self.errori.append(f"Lettura non riuscita di {nome}: {e}")
                    continue
                if raw.startswith(b"\x89PNG") or raw.startswith(b"\xff\xd8\xff") or raw.startswith(b"GIF8"):
                    self.immagini.append(nome)
                    continue
                if estensione == ".xml" or raw.lstrip()[:5] == b"<?xml":
                    radice = self.parse_xml(raw)
                    if radice is not None:
                        self.xml[nome] = radice
                        continue
                testo = decodifica_bytes(raw)
                if testo is None:
                    continue
                if testo.lstrip().startswith("<?xml"):
                    radice = self.parse_xml(testo)
                    if radice is not None:
                        self.xml[nome] = radice
                        continue
                self.testi[nome] = testo
                normalizzato = normalizza_percorso(nome)
                self.testi_normalizzati[normalizzato] = testo
                self.testi_normalizzati.setdefault(normalizzato.split("/")[-1], testo)
                base = nome.replace("\\", "/").split("/")[-1]
                if not base.isdigit():
                    self.risorse_per_nome.setdefault(base, nome)

    def indicizza(self):
        for nome_file, radice in self.xml.items():
            indice = {"file": nome_file, "tipi": {}, "stringhe": {}, "oggetti": {}, "genitori": {}}
            for padre in radice.iter():
                for figlio in padre:
                    indice["genitori"][figlio] = padre
            for elemento in radice.iter():
                tag = elemento.tag
                if not isinstance(tag, str):
                    continue
                ident = elemento.get("Id")
                if tag == "TypeObject" and ident is not None:
                    indice["tipi"][ident] = elemento.get("FullTypeName", "")
                elif tag == "Object" and ident is not None:
                    indice["oggetti"][ident] = elemento
                elif tag == "String" and ident is not None:
                    indice["stringhe"][ident] = elemento.get("Value", elemento.text or "")
                percorso = elemento.get("ArchiveElementPath") or elemento.get("Path")
                if "resource" in tag.lower() and percorso:
                    ident_risorsa = elemento.get("Id") or elemento.get("ResourceId") or ""
                    nome_risorsa = elemento.get("Name", "")
                    self.risorse.append({"id": ident_risorsa, "nome": nome_risorsa, "percorso": percorso})
                    if ident_risorsa:
                        self.risorse_per_id[ident_risorsa] = percorso
                    for chiave in (nome_risorsa, percorso, percorso.replace("\\", "/").split("/")[-1]):
                        if chiave and not chiave.isdigit():
                            self.risorse_per_nome[chiave] = percorso
            self.indici[nome_file] = indice
        documenti = []
        for nome_file, radice in self.xml.items():
            indice = self.indici[nome_file]
            indice["copia"] = False
            for oggetto in radice.iter("Object"):
                if tipo_base(self.tipo_oggetto(oggetto, indice)) == "Spotfire.Dxp.Application.Document":
                    documenti.append(nome_file)
                    break
        principale = next((f for f in documenti if f.lower().endswith("analysisdocument.xml")), documenti[0] if documenti else None)
        for nome_file in documenti:
            if nome_file != principale:
                self.indici[nome_file]["copia"] = True
                self.copie_documento.append(nome_file)
        self.file_principale = principale
        self.nome_risorsa_per_file = {normalizza_percorso(r["percorso"]): r["nome"] for r in self.risorse}
        for nome_file, radice in self.xml.items():
            indice = self.indici[nome_file]
            for elemento in radice.iter():
                if not isinstance(elemento.tag, str):
                    continue
                if "resource" in elemento.tag.lower() and (elemento.get("ArchiveElementPath") or elemento.get("Path")):
                    continue
                coppie = list(elemento.attrib.items())
                if elemento.text and len(elemento.text) < 300:
                    coppie.append(("", elemento.text.strip()))
                for chiave, valore in coppie:
                    percorso = self.risorsa_da_valore(valore, chiave, elemento.tag)
                    if percorso:
                        self.riferimenti_risorse.setdefault(normalizza_percorso(percorso), []).append((elemento, indice))

    def risorsa_da_valore(self, valore, chiave, tag):
        if not valore:
            return None
        if valore in self.risorse_per_nome:
            return self.risorse_per_nome[valore]
        indizio = f"{chiave}{tag}".lower()
        if valore in self.risorse_per_id and ("resource" in indizio or "embedded" in indizio or "blob" in indizio):
            return self.risorse_per_id[valore]
        return None

    def testo_risorsa(self, percorso):
        if percorso in self.testi:
            return self.testi[percorso]
        normalizzato = normalizza_percorso(percorso)
        return self.testi_normalizzati.get(normalizzato) or self.testi_normalizzati.get(normalizzato.split("/")[-1])

    def tipo_oggetto(self, oggetto, indice):
        if oggetto in self.cache_tipo:
            return self.cache_tipo[oggetto]
        risultato = ""
        tipo = oggetto.find("Type")
        if tipo is None:
            risultato = oggetto.get("FullTypeName", "")
        else:
            for figlio in tipo:
                if figlio.tag == "TypeObject":
                    risultato = figlio.get("FullTypeName", "")
                    break
                if isinstance(figlio.tag, str) and figlio.tag.endswith("Ref"):
                    risultato = indice["tipi"].get(figlio.get("Value") or figlio.get("Id"), "")
                    break
            if not risultato:
                for valore in tipo.attrib.values():
                    if valore in indice["tipi"]:
                        risultato = indice["tipi"][valore]
                        break
        self.cache_tipo[oggetto] = risultato
        return risultato

    def in_istantanea(self, elemento, indice):
        if indice.get("copia"):
            return True
        genitori = indice["genitori"]
        nodo = genitori.get(elemento)
        catena = []
        esito = False
        while nodo is not None:
            if nodo in self.cache_istantanea:
                esito = self.cache_istantanea[nodo]
                break
            catena.append(nodo)
            if nodo.tag == "Object" and TIPI_ISTANTANEA.search(ultimo_segmento(self.tipo_oggetto(nodo, indice))):
                esito = True
                break
            nodo = genitori.get(nodo)
        for n in catena:
            self.cache_istantanea[n] = esito
        return esito

    def antenati_oggetti(self, elemento, indice, includi_se_stesso=True):
        nodo = elemento if includi_se_stesso else indice["genitori"].get(elemento)
        while nodo is not None:
            if nodo.tag == "Object":
                yield nodo
            nodo = indice["genitori"].get(nodo)

    def catena_tipi(self, elemento, indice, massimo=6):
        tipi = [ultimo_segmento(self.tipo_oggetto(o, indice)) for o in self.antenati_oggetti(elemento, indice)]
        return " dentro ".join(t for t in tipi[:massimo] if t)

    def risolvi(self, elemento, indice):
        if elemento.tag == "Object":
            return elemento
        if isinstance(elemento.tag, str) and elemento.tag.endswith("Ref") and elemento.tag not in ("TypeRef", "StringRef"):
            return indice["oggetti"].get(elemento.get("Value") or elemento.get("Id"))
        return None

    def valore_semplice(self, campo, indice, profondita=2):
        for figlio in campo:
            if figlio.tag == "String":
                valore = figlio.get("Value", figlio.text or "")
                if valore.strip():
                    return valore
            elif isinstance(figlio.tag, str) and figlio.tag.endswith("Ref"):
                valore = indice["stringhe"].get(figlio.get("Value") or figlio.get("Id"))
                if valore:
                    return valore
            elif figlio.tag == "Object" and profondita > 0:
                campi = figlio.find("Fields")
                if campi is None:
                    continue
                for interno in campi:
                    if chiave_campo(interno.get("Name")) in ("expression", "value", "text", "name", "title"):
                        valore = self.valore_semplice(interno, indice, profondita - 1)
                        if valore:
                            return valore
        return ""

    def nome_diretto(self, oggetto, indice):
        campi = oggetto.find("Fields")
        if campi is None:
            return ""
        trovati = {}
        for campo in campi:
            if campo.tag != "Field":
                continue
            chiave = chiave_campo(campo.get("Name"))
            if chiave in CHIAVI_NOME and chiave not in trovati:
                valore = self.valore_semplice(campo, indice)
                if valore:
                    trovati[chiave] = valore
        for chiave in CHIAVI_NOME:
            if chiave in trovati:
                return trovati[chiave].strip()
        return ""

    def raccogli_valori(self, elemento, indice, percorso="", profondita=0):
        risultati = []
        if profondita > 150:
            return risultati
        for figlio in elemento:
            tag = figlio.tag
            if not isinstance(tag, str) or tag == "Type":
                continue
            if tag == "Field":
                nome = figlio.get("Name", "")
                sotto = f"{percorso}.{nome}" if percorso else nome
                risultati.extend(self.raccogli_valori(figlio, indice, sotto, profondita + 1))
                continue
            if tag.endswith("Ref"):
                riferimento = figlio.get("Value") or figlio.get("Id")
                if riferimento in indice["stringhe"]:
                    risultati.append((percorso, indice["stringhe"][riferimento]))
                continue
            valore = figlio.get("Value")
            if valore is None and len(figlio) == 0 and figlio.text and figlio.text.strip():
                valore = figlio.text
            if valore is not None:
                risultati.append((percorso, valore))
            for chiave, v in figlio.attrib.items():
                percorso_risorsa = self.risorsa_da_valore(v, chiave, tag)
                if percorso_risorsa:
                    contenuto = self.testo_risorsa(percorso_risorsa)
                    if contenuto:
                        risultati.append((percorso, contenuto))
            risultati.extend(self.raccogli_valori(figlio, indice, percorso, profondita + 1))
        return risultati

    @staticmethod
    def valore_campo(valori, chiavi, modo="ultimo", escludi=()):
        for percorso, valore in valori:
            minuscolo = percorso.lower()
            if any(e in minuscolo for e in escludi):
                continue
            parti = [chiave_campo(p) for p in minuscolo.split(".") if p]
            if not parti:
                continue
            corrisponde = parti[-1] in chiavi if modo == "ultimo" else any(p in chiavi for p in parti)
            if corrisponde and valore and valore.strip():
                return valore
        return ""

    def censisci_oggetti(self):
        for nome_file, radice in self.xml.items():
            indice = self.indici[nome_file]
            if indice.get("copia"):
                continue
            for oggetto in radice.iter("Object"):
                tipo = self.tipo_oggetto(oggetto, indice)
                if not tipo:
                    continue
                istantanea = self.in_istantanea(oggetto, indice)
                if istantanea:
                    self.tipi_istantanea[tipo] += 1
                else:
                    self.tipi_vivi[tipo] += 1
                ultimo = ultimo_segmento(tipo)
                if ultimo == "Bookmark":
                    self.segnalibri += 1
                if istantanea or "+" in tipo_base(tipo):
                    continue
                if ultimo.endswith("ScriptDefinition"):
                    self.oggetti_script.append((oggetto, indice))
                elif ultimo.endswith("DataFunctionDefinition"):
                    self.oggetti_data_function.append((oggetto, indice))
                elif re.fullmatch(r"DataProperty(?:Impl)?", ultimo):
                    self.oggetti_proprieta.append((oggetto, indice))
                elif ultimo in ("CalculatedColumn", "CalculatedColumnImpl"):
                    self.oggetti_colonne_calcolate.append((oggetto, indice))

    def elementi_collezione(self, oggetto, indice, nome_campo):
        campi = oggetto.find("Fields")
        if campi is None:
            return None
        campo = next((c for c in campi if c.tag == "Field" and chiave_campo(c.get("Name")) == nome_campo), None)
        if campo is None:
            return None
        contenitore = campo
        for figlio in campo:
            risolto = self.risolvi(figlio, indice)
            if risolto is not None:
                contenitore = risolto
                break
        elementi = None
        for interno in contenitore.iter("Field"):
            if chiave_campo(interno.get("Name")) == "items":
                elementi = interno.find(".//Elements")
                break
        if elementi is None:
            elementi = contenitore.find(".//Elements")
        if elementi is None:
            return None
        return [c for c in elementi if self.risolvi(c, indice) is not None]

    def visual_di_pagina(self, pagina, indice, visti, segui_riferimenti):
        risultati = []
        pila = list(reversed(list(pagina)))
        while pila:
            nodo = pila.pop()
            if not isinstance(nodo.tag, str) or nodo.tag == "Type":
                continue
            bersaglio = self.risolvi(nodo, indice)
            if bersaglio is not None:
                if nodo.tag != "Object" and not segui_riferimenti:
                    continue
                tipo = self.tipo_oggetto(bersaglio, indice)
                if e_visual(tipo):
                    if bersaglio not in visti:
                        visti.add(bersaglio)
                        risultati.append((ultimo_segmento(tipo), self.nome_diretto(bersaglio, indice), bersaglio))
                    continue
                if nodo.tag != "Object":
                    continue
            pila.extend(reversed(list(nodo)))
        return risultati

    def primo_visual_interno(self, oggetto, indice):
        pila = list(reversed(list(oggetto)))
        generico = None
        while pila:
            nodo = pila.pop()
            if not isinstance(nodo.tag, str) or nodo.tag == "Type":
                continue
            if nodo.tag == "Object":
                tipo = self.tipo_oggetto(nodo, indice)
                if e_visual(tipo):
                    return nodo
                base = tipo_base(tipo)
                if generico is None and "+" not in base and base.startswith("Spotfire.Dxp.Application.Visuals.") and "Collection" not in base:
                    generico = nodo
            pila.extend(reversed(list(nodo)))
        return generico

    def visual_da_collezione(self, pagina, indice, visti):
        elementi = self.elementi_collezione(pagina, indice, "visuals")
        if elementi is None:
            return None
        risultati = []
        for elemento in elementi:
            oggetto = self.risolvi(elemento, indice)
            if oggetto is None or oggetto in visti:
                continue
            visti.add(oggetto)
            bersaglio = oggetto
            if not e_visual(self.tipo_oggetto(oggetto, indice)):
                interno = self.primo_visual_interno(oggetto, indice)
                if interno is not None:
                    bersaglio = interno
            tipo = ultimo_segmento(self.tipo_oggetto(bersaglio, indice)) or "Sconosciuto"
            nome = self.nome_diretto(bersaglio, indice) or self.nome_diretto(oggetto, indice)
            risultati.append((tipo, nome, bersaglio))
        return risultati

    def estrai_pagine_e_visual(self):
        documento = None
        file_ordinati = sorted(self.xml.items(), key=lambda x: 0 if x[0] == self.file_principale else 1)
        for nome_file, radice in file_ordinati:
            indice = self.indici[nome_file]
            if indice.get("copia"):
                continue
            for oggetto in radice.iter("Object"):
                if tipo_base(self.tipo_oggetto(oggetto, indice)) == "Spotfire.Dxp.Application.Document" and not self.in_istantanea(oggetto, indice):
                    documento = (oggetto, indice)
                    break
            if documento:
                break
        pagine = []
        if documento:
            oggetto, indice = documento
            for elemento in self.elementi_collezione(oggetto, indice, "pages") or []:
                pagina = self.risolvi(elemento, indice)
                if pagina is not None and ultimo_segmento(self.tipo_oggetto(pagina, indice)) == "Page" and pagina not in [p for p, _ in pagine]:
                    pagine.append((pagina, indice))
        if not pagine:
            for nome_file, radice in self.xml.items():
                indice = self.indici[nome_file]
                if indice.get("copia"):
                    continue
                for oggetto in radice.iter("Object"):
                    if tipo_base(self.tipo_oggetto(oggetto, indice)) == "Spotfire.Dxp.Application.Page" and not self.in_istantanea(oggetto, indice):
                        pagine.append((oggetto, indice))
        visti = set()
        for numero, (pagina, indice) in enumerate(pagine, 1):
            titolo = self.nome_diretto(pagina, indice) or f"Pagina {numero}"
            visual = self.visual_da_collezione(pagina, indice, visti)
            if visual is None:
                visual = self.visual_di_pagina(pagina, indice, visti, False)
            self.pagine.append({"titolo": titolo, "visual": [(u, n) for u, n, _ in visual]})
            for ultimo, nome, oggetto in visual:
                self.visual[ultimo] += 1
                if ultimo in TIPI_TEXT_AREA:
                    self.aggiungi_text_area(oggetto, indice, nome or f"Text area {len(self.text_area) + 1}", titolo)

    def aggiungi_text_area(self, oggetto, indice, titolo, pagina):
        valori = self.raccogli_valori(oggetto, indice)
        preferiti = [v for p, v in valori if "html" in p.lower() and v.strip()]
        generici = [v for p, v in valori if "<" in v and ">" in v and len(v) > 5]
        candidati = preferiti or generici
        html = max(candidati, key=len) if candidati else ""
        self.text_area.append({
            "titolo": titolo,
            "pagina": pagina,
            "html": html,
            "controlli": len(re.findall(r"<SpotfireControl\b", html, re.I)),
            "immagini": len(re.findall(r"<img\b", html, re.I)),
            "tag_script": len(re.findall(r"<script\b", html, re.I)),
            "librerie": re.findall(r"<script\b[^>]*\bsrc=[\"']([^\"']+)", html, re.I),
        })

    def scegli_codice(self, valori, anche_senza_segnali=False):
        migliore = ""
        punteggio_migliore = (-1, 0)
        for percorso, valore in valori:
            if not valore or not valore.strip():
                continue
            chiave = chiave_campo(percorso.split(".")[-1]) if percorso else ""
            if not (chiave in CHIAVI_CODICE or chiave.endswith("code")):
                continue
            segnali = max(punti_codice(valore))
            if segnali == 0 and not (anche_senza_segnali and "\n" in valore):
                continue
            punteggio = (segnali, len(valore))
            if punteggio > punteggio_migliore:
                migliore, punteggio_migliore = valore, punteggio
        return migliore

    def estrai_data_function(self):
        for oggetto, indice in self.oggetti_data_function:
            valori = self.raccogli_valori(oggetto, indice)
            codice = self.scegli_codice(valori, anche_senza_segnali=True)
            dichiarato = self.valore_campo(valori, {"language", "scriptlanguage", "languagename", "engine", "executor"}, modo="qualsiasi").lower()
            if "python" in dichiarato or re.search(r"^\s*(?:import|from|def)\s", codice, re.M):
                linguaggio = "Python"
            else:
                linguaggio = "R (TERR)"
            if codice:
                self.hash_data_function.add(impronta(codice))
            righe_codice = len([r for r in codice.splitlines() if r.strip()])
            self.data_function.append({
                "nome": self.nome_diretto(oggetto, indice) or f"Data function {len(self.data_function) + 1}",
                "linguaggio": linguaggio,
                "codice": codice,
                "righe": righe_codice,
                "ingressi": sorted({v for p, v in valori if re.search(r"input", p, re.I) and chiave_campo(p.split(".")[-1]) == "name"}),
                "uscite": sorted({v for p, v in valori if re.search(r"output", p, re.I) and chiave_campo(p.split(".")[-1]) == "name"}),
            })

    def nome_elemento(self, nodo, indice):
        if nodo.tag not in ("Field", "Attribute", "Property", "String"):
            for chiave in ("Name", "name", "DisplayName", "Title", "ScriptName"):
                valore = nodo.get(chiave)
                if valore and valore.strip() and len(valore) < 200:
                    return valore.strip()
        if nodo.tag == "Object":
            return self.nome_diretto(nodo, indice)
        for figlio in nodo:
            if not isinstance(figlio.tag, str):
                continue
            etichetta = figlio.tag.lower()
            if etichetta in ("name", "title", "displayname", "scriptname") and figlio.text and figlio.text.strip() and len(figlio.text) < 200:
                return figlio.text.strip()
            if figlio.tag == "Field" and chiave_campo(figlio.get("Name")) in ("name", "title", "displayname", "scriptname"):
                valore = self.valore_semplice(figlio, indice)
                if valore and len(valore) < 200:
                    return valore.strip()
            if figlio.tag in ("Attribute", "Property") and (figlio.get("Key") or figlio.get("Name") or "").lower() in ("name", "title", "displayname", "scriptname"):
                valore = figlio.get("Value") or figlio.text or ""
                if valore.strip() and len(valore) < 200:
                    return valore.strip()
        return ""

    def linguaggio_elemento(self, nodo, indice):
        for chiave, valore in nodo.attrib.items():
            if "language" in chiave.lower():
                return valore
        for figlio in nodo:
            if not isinstance(figlio.tag, str):
                continue
            if "language" in figlio.tag.lower() and figlio.text:
                return figlio.text
            etichetta = (figlio.get("Name") or figlio.get("Key") or "").lower() if figlio.tag in ("Field", "Attribute", "Property") else ""
            if "language" in etichetta:
                valore = figlio.get("Value") or self.valore_semplice(figlio, indice, 3)
                for p, v in self.raccogli_valori(figlio, indice)[:5]:
                    valore = valore or v
                return valore or ""
        return ""

    def contesto_vicino(self, elemento, indice, contatore):
        nodo = elemento
        livelli = 0
        nome = ""
        linguaggio = ""
        while nodo is not None and livelli < 7:
            if nodo is not elemento and contatore.get(nodo, 0) > 1:
                break
            nome = nome or self.nome_elemento(nodo, indice)
            linguaggio = linguaggio or self.linguaggio_elemento(nodo, indice)
            if nome and linguaggio:
                break
            nodo = indice["genitori"].get(nodo)
            livelli += 1
        return nome, linguaggio

    def estrai_script(self):
        file_script = set()
        for nome_file in self.xml:
            nome_risorsa = self.nome_risorsa_per_file.get(normalizza_percorso(nome_file), "")
            if "script" in nome_risorsa.lower() or "script" in Path(nome_file).name.lower():
                file_script.add(nome_file)

        trovati = []
        visti = set(self.hash_data_function)
        for nome_file, radice in self.xml.items():
            indice = self.indici[nome_file]
            if indice.get("copia"):
                continue
            dedicato = nome_file in file_script
            for elemento in radice.iter():
                for valore in (elemento.get("Value"), elemento.text):
                    if not valore or not valore.strip() or valore.lstrip().startswith("<"):
                        continue
                    if dedicato:
                        if len(valore.strip()) < 8 or max(punti_codice(valore)) < 1:
                            continue
                        linguaggio = indovina_linguaggio(valore) or indovina_linguaggio(valore, rigoroso=False)
                    else:
                        if len(valore.strip()) < 15:
                            continue
                        linguaggio = indovina_linguaggio(valore)
                    if not linguaggio:
                        continue
                    chiave = impronta(valore)
                    if chiave in visti:
                        continue
                    if self.in_istantanea(elemento, indice):
                        continue
                    visti.add(chiave)
                    trovati.append({"testo": valore, "linguaggio": linguaggio, "elemento": elemento, "indice": indice, "dedicato": dedicato, "origine": self.nome_risorsa_per_file.get(normalizza_percorso(nome_file)) or nome_file})
        for percorso, testo in self.testi.items():
            linguaggio = indovina_linguaggio(testo)
            chiave = impronta(testo)
            if not linguaggio or chiave in visti:
                continue
            normalizzato = normalizza_percorso(percorso)
            riferimenti = self.riferimenti_risorse.get(normalizzato, []) + self.riferimenti_risorse.get(normalizzato.split("/")[-1], [])
            if riferimenti and all(self.in_istantanea(e, i) for e, i in riferimenti):
                continue
            visti.add(chiave)
            trovati.append({"testo": testo, "linguaggio": linguaggio, "elemento": None, "indice": None, "dedicato": False, "origine": self.nome_risorsa_per_file.get(normalizzato) or percorso})
        trovati.sort(key=lambda t: 0 if t["dedicato"] else 1)

        contatori = {}
        for t in trovati:
            if t["elemento"] is None:
                continue
            contatore = contatori.setdefault(t["indice"]["file"], Counter())
            nodo = t["elemento"]
            while nodo is not None:
                contatore[nodo] += 1
                nodo = t["indice"]["genitori"].get(nodo)

        nomi_visti = set()
        for t in trovati:
            nome = ""
            dichiarato = ""
            contesto = ""
            if t["elemento"] is not None:
                nome, dichiarato = self.contesto_vicino(t["elemento"], t["indice"], contatori[t["indice"]["file"]])
                contesto = self.catena_tipi(t["elemento"], t["indice"])
            if nome and impronta(nome) == impronta(t["testo"]):
                nome = ""
            linguaggio = linguaggio_dichiarato(dichiarato) or t["linguaggio"]
            chiave_nome = (linguaggio, nome.lower())
            if nome and chiave_nome in nomi_visti:
                continue
            nomi_visti.add(chiave_nome)
            self.script.append({
                "nome": nome,
                "nome_nel_file": bool(nome),
                "linguaggio": linguaggio,
                "codice": t["testo"],
                "righe": len([r for r in t["testo"].splitlines() if r.strip()]),
                "origine": t["origine"],
            })

        progressivi = Counter()
        for s in self.script:
            progressivi[s["linguaggio"]] += 1
            if not s["nome"]:
                s["nome"] = f"Script {s['linguaggio']} {progressivi[s['linguaggio']]}"

    def classe_proprieta(self, oggetto, indice):
        nodo = indice["genitori"].get(oggetto)
        passi = 0
        while nodo is not None and passi < 30:
            if nodo.tag == "Field":
                nome = chiave_campo(nodo.get("Name"))
                if nome != "properties" and ("properties" in nome or re.search(r"document|column|table|hierarchy|filter|visual", nome)):
                    return nodo.get("Name")
            nodo = indice["genitori"].get(nodo)
            passi += 1
        return ""

    @staticmethod
    def visibile_in_properties(proprieta, con_attributi):
        if not proprieta["standard"]:
            return True
        if con_attributi and "isvisible" not in proprieta["attributi"].lower():
            return False
        if "." in proprieta["nome"]:
            return False
        return proprieta["nome"].lower() not in PROPRIETA_SCHEDA_GENERAL

    def estrai_proprieta(self):
        citate = {}
        for s in self.script:
            for nome in re.findall(r"Properties\s*\[\s*[\"']([^\"']+)[\"']\s*\]", s["codice"]):
                citate.setdefault(nome, set()).add(s["nome"])
            for nome in re.findall(r"Properties\.\w+\(\s*[\"']([^\"']+)[\"']", s["codice"]):
                citate.setdefault(nome, set()).add(s["nome"])
        for nome_file, radice in self.xml.items():
            indice = self.indici[nome_file]
            if indice.get("copia"):
                continue
            for elemento in radice.iter():
                valore = elemento.get("Value")
                if not valore or ("${" not in valore and "DocumentProperty" not in valore):
                    continue
                nomi = re.findall(r"\$\{([^}]+)\}", valore) + re.findall(r"DocumentProperty\(\s*\"([^\"]+)\"", valore)
                if nomi and not self.in_istantanea(elemento, indice):
                    for nome in nomi:
                        citate.setdefault(nome.strip(), set()).add("espressioni")
        self.proprieta_citate = citate

        viste = set()
        for oggetto, indice in self.oggetti_proprieta:
            valori = self.raccogli_valori(oggetto, indice)
            nome = self.nome_diretto(oggetto, indice)
            if not nome:
                continue
            classe = self.classe_proprieta(oggetto, indice)
            chiave = (classe.lower(), nome)
            if chiave in viste:
                continue
            viste.add(chiave)
            attributi = self.valore_campo(valori, {"attributes"})
            minuscola = classe.lower()
            if "document" in minuscola or "analysis" in minuscola:
                gruppo = "Documento"
            elif "column" in minuscola:
                gruppo = "Colonna"
            elif "table" in minuscola:
                gruppo = "Tabella"
            else:
                gruppo = classe or "Non determinata"
            self.proprieta.append({
                "nome": nome,
                "classe": classe,
                "gruppo": gruppo,
                "standard": "isstandard" in attributi.lower(),
                "attributi": attributi,
                "tipo_dato": self.valore_campo(valori, {"datatype", "valuetype", "typename"}, modo="qualsiasi", escludi=("description",)),
                "valore": self.valore_campo(valori, {"value", "defaultvalue", "currentvalue"}, escludi=("datatype", "description")),
                "descrizione": self.valore_campo(valori, {"description"}),
                "usata_in": ", ".join(sorted(citate.get(nome, []))),
            })

        documento = [p for p in self.proprieta if p["gruppo"] == "Documento"]
        if documento:
            con_attributi = any(p["attributi"] for p in documento)
            visibili = [p for p in documento if self.visibile_in_properties(p, con_attributi)]
            self.proprieta_documento = sorted(visibili, key=lambda p: (not p["standard"],))
            self.proprieta_documento_standard = [p for p in self.proprieta_documento if p["standard"]]
        else:
            non_standard = [p for p in self.proprieta if not p["standard"] and p["gruppo"] not in ("Colonna", "Tabella")]
            if non_standard:
                self.proprieta_documento = non_standard
            else:
                for nome, fonti in sorted(citate.items()):
                    self.proprieta_documento.append({"nome": nome, "classe": "", "gruppo": "", "standard": False, "attributi": "", "tipo_dato": "", "valore": "", "descrizione": "", "usata_in": ", ".join(sorted(fonti))})

    def estrai_tabelle_sorgente(self):
        chiavi_utili = re.compile(r"catalog|schema|table|view|owner|database|external|tabletype|elementid|librarypath|path|server|connection|sourcetype|customquery|query", re.I)
        viste = set()
        for nome_file, radice in self.xml.items():
            indice = self.indici[nome_file]
            if indice.get("copia"):
                continue
            for elemento in radice.iter():
                if not isinstance(elemento.tag, str) or not re.search(r"TableSchema$|^DataTableSource$|^SourceTable$", elemento.tag):
                    continue
                nome = elemento.get("DisplayName") or elemento.get("Name") or ""
                proprieta = OrderedDict()
                for k, v in elemento.attrib.items():
                    if k not in ("Name", "DisplayName") and chiavi_utili.search(k):
                        proprieta[k] = v
                pila = list(elemento)
                while pila:
                    figlio = pila.pop(0)
                    if not isinstance(figlio.tag, str) or "column" in figlio.tag.lower():
                        continue
                    chiave = figlio.get("Key") or (figlio.get("Name") if figlio.tag in ("Attribute", "Property", "Field") else None)
                    valore = figlio.get("Value") if figlio.get("Value") is not None else (figlio.text or "").strip()
                    if chiave and valore and chiavi_utili.search(chiave) and chiave not in proprieta:
                        proprieta[chiave] = valore
                    pila.extend(list(figlio))
                self.aggiungi_tabella(nome, proprieta, nome_file)
            for oggetto in radice.iter("Object"):
                tipo = self.tipo_oggetto(oggetto, indice)
                ultimo = ultimo_segmento(tipo)
                if not re.search(r"DataSource$|InformationLink|DatabaseTable|SourceTable", ultimo) or re.search(r"Collection|Settings|Manager", ultimo):
                    continue
                if self.in_istantanea(oggetto, indice):
                    continue
                proprieta = OrderedDict([("TipoSorgente", ultimo)])
                for percorso, valore in self.raccogli_valori(oggetto, indice)[:400]:
                    ultima = percorso.split(".")[-1] if percorso else ""
                    if ultima and valore and len(valore) < 2000 and chiavi_utili.search(ultima) and ultima not in proprieta:
                        proprieta[ultima] = valore
                if len(proprieta) > 1:
                    self.aggiungi_tabella(self.nome_diretto(oggetto, indice), proprieta, nome_file)

    def aggiungi_tabella(self, nome, proprieta, nome_file):
        chiave = nome.strip().lower() if nome else tuple(proprieta.items())
        if chiave in self.tabelle_viste or (not nome and not proprieta):
            return
        self.tabelle_viste.add(chiave)
        valori = {k.lower(): v for k, v in proprieta.items()}
        query = next((v for k, v in valori.items() if "customquery" in k or k in ("query", "sqlquery", "sql")), "")
        catalogo = next((v for k, v in valori.items() if "catalog" in k and "path" not in k), "")
        schema = next((v for k, v in valori.items() if "schema" in k and "path" not in k), "")
        oggetto_db = next((v for k, v in valori.items() if re.search(r"^(?:table|tablename|view|viewname|sourcetable|sourcetablename|externalname|externaltablename)$", k)), "")
        tipo_oggetto_db = next((v for k, v in valori.items() if "tabletype" in k or k == "type"), "")
        if query:
            natura = "Query personalizzata"
        elif oggetto_db or schema:
            natura = "Tabella o vista del database"
        elif any("librarypath" in k or "elementid" in k for k in valori) or "InformationLink" in proprieta.get("TipoSorgente", ""):
            natura = "Information Link"
        else:
            natura = "Non determinato"
        completo = ".".join(x for x in (catalogo, schema, oggetto_db or (nome if natura == "Tabella o vista del database" else "")) if x)
        self.tabelle_sorgente.append({
            "nome": nome,
            "natura": natura,
            "oggetto_database": completo,
            "tipo_oggetto_database": tipo_oggetto_db,
            "dettagli": "; ".join(f"{k}={re.sub(chr(10), ' ', v)[:150]}" for k, v in proprieta.items() if "query" not in k.lower()),
            "file": nome_file,
        })

    def estrai_query_e_connessioni(self):
        file_ordinati = sorted(self.xml.items(), key=lambda x: 0 if x[0].lower().endswith("dataaccessplan.xml") else 1)
        for nome_file, radice in file_ordinati:
            if self.indici[nome_file].get("copia"):
                continue
            pila = [(radice, "")]
            while pila:
                elemento, contesto = pila.pop()
                if not isinstance(elemento.tag, str):
                    continue
                tag = elemento.tag
                nome_elemento = elemento.get("DisplayName") or elemento.get("Name")
                if tag not in ("Field", "Attribute", "Property") and nome_elemento:
                    contesto = nome_elemento
                chiave = elemento.get("Key") or (elemento.get("Name") if tag in ("Field", "Property") else None)
                if chiave:
                    valore = elemento.get("Value")
                    if valore is None:
                        valori_figli = [f.get("Value") or (f.text or "") for f in elemento]
                        valore = max(valori_figli, key=len) if valori_figli else (elemento.text or "")
                    if valore:
                        self.valuta_coppia(chiave, valore, contesto, nome_file)
                for k, v in elemento.attrib.items():
                    if k not in ("Key", "Name", "Value", "DisplayName", "Id"):
                        self.valuta_coppia(k, v, contesto, nome_file)
                for figlio in elemento:
                    pila.append((figlio, contesto))

    def valuta_coppia(self, chiave, valore, contesto, nome_file):
        minuscola = chiave.lower()
        if re.search(r"pass|pwd|secret|token|credential", minuscola):
            return
        if re.search(r"query|sql|commandtext|statement", minuscola) and len(valore) > 15 and re.search(r"\b(?:select|with|exec|execute|call)\b", valore, re.I):
            chiave_hash = impronta(valore)
            chiave_nome = (contesto or "").strip().lower()
            if chiave_hash in self.hash_query or (chiave_nome and chiave_nome in self.nomi_query):
                return
            self.hash_query.add(chiave_hash)
            if chiave_nome:
                self.nomi_query.add(chiave_nome)
            self.query.append({"nome": contesto or f"query {len(self.query) + 1}", "sql": valore, "file": ""})
            return
        if re.search(r"server|host|database|catalog|schema|adaptertype|connectionstring|datasourcename|httppath|warehouse|driver|provider|librarypath|informationlink", minuscola) and 0 < len(valore) < 500:
            valore_sicuro = re.sub(r"(?i)(password|pwd)\s*=\s*[^;]*", r"\1=nascosta", valore)
            coppia = (contesto, chiave, valore_sicuro)
            if coppia not in self.connessioni_viste:
                self.connessioni_viste.add(coppia)
                self.connessioni.append({"contesto": contesto, "chiave": chiave, "valore": valore_sicuro})

    def estrai_colonne_calcolate(self):
        visti = set()
        for oggetto, indice in self.oggetti_colonne_calcolate:
            valori = self.raccogli_valori(oggetto, indice)
            nome = ""
            tabella = ""
            for antenato in self.antenati_oggetti(oggetto, indice):
                ultimo = ultimo_segmento(self.tipo_oggetto(antenato, indice))
                if not nome:
                    nome = self.nome_diretto(antenato, indice)
                if ultimo == "DataTable":
                    tabella = self.nome_diretto(antenato, indice)
                    break
            espressione = self.valore_campo(valori, {"expression", "expressiontext", "originalexpression"}, modo="qualsiasi")
            chiave = (tabella, nome, espressione)
            if chiave in visti:
                continue
            visti.add(chiave)
            self.colonne_calcolate.append({"tabella": tabella, "nome": nome, "espressione": espressione})

    def estrai_archivio(self, destinazione):
        with zipfile.ZipFile(self.percorso, "r") as zf:
            for info in zf.infolist():
                if info.is_dir():
                    continue
                parti = [pulisci_nome(p, 120) for p in re.split(r"[\\/]", info.filename) if p and p not in (".", "..")]
                if not parti:
                    continue
                obiettivo = destinazione.joinpath(*parti)
                try:
                    obiettivo.parent.mkdir(parents=True, exist_ok=True)
                    with zf.open(info) as sorgente, open(obiettivo, "wb") as uscita:
                        uscita.write(sorgente.read())
                except OSError as e:
                    self.errori.append(f"{info.filename}: {e}")

    def esporta(self, cartella):
        if cartella.exists():
            shutil.rmtree(cartella, ignore_errors=True)
            if cartella.exists() and any(p.is_file() for p in cartella.rglob("*")):
                self.avvisi.append(f"file non eliminati in {cartella}, chiudere i file aperti e rilanciare")
        cartella.mkdir(parents=True, exist_ok=True)
        if ESTRAI_ARCHIVIO_COMPLETO:
            self.estrai_archivio(cartella / "archivio_estratto")

        sottocartelle = {"IronPython": ("ironpython", ".py"), "Python": ("python", ".py"), "JavaScript": ("javascript", ".js")}
        usati = {}
        for s in self.script:
            sotto, estensione = sottocartelle.get(s["linguaggio"], ("altri_script", ".txt"))
            nome_file = nome_file_libero(s["nome"], estensione, usati.setdefault(sotto, set()))
            s["file"] = f"{sotto}/{nome_file}"
            scrivi_testo(cartella / sotto / nome_file, s["codice"], self.errori)

        usati_df = set()
        for d in self.data_function:
            estensione = ".py" if d["linguaggio"] == "Python" else ".R"
            nome_file = nome_file_libero(d["nome"], estensione, usati_df)
            d["file"] = f"data_function/{nome_file}" if d["codice"] else ""
            if d["codice"]:
                scrivi_testo(cartella / "data_function" / nome_file, d["codice"], self.errori)

        usati_html = set()
        for t in self.text_area:
            t["file"] = ""
            if t["html"].strip():
                nome_file = nome_file_libero(f"{t['pagina']} - {t['titolo']}", ".html", usati_html)
                t["file"] = f"text_area_html/{nome_file}"
                scrivi_testo(cartella / "text_area_html" / nome_file, t["html"], self.errori)

        usati_sql = set()
        for q in self.query:
            nome_file = nome_file_libero(q["nome"], ".sql", usati_sql)
            q["file"] = f"query_sql/{nome_file}"
            scrivi_testo(cartella / "query_sql" / nome_file, q["sql"], self.errori)

        scrivi_tabella(
            cartella / "script",
            ["nome", "linguaggio", "file", "righe_di_codice", "nome_presente_nel_file", "risorsa_origine"],
            [[s["nome"], s["linguaggio"], s["file"], s["righe"], "si" if s["nome_nel_file"] else "no", s["origine"]] for s in self.script],
        )
        scrivi_tabella(
            cartella / "pagine",
            ["numero", "pagina", "oggetti_visivi", "text_area", "grafici_e_tabelle"],
            [[i, p["titolo"], len(p["visual"]), sum(1 for u, _ in p["visual"] if u in TIPI_TEXT_AREA), sum(1 for u, _ in p["visual"] if u not in TIPI_TEXT_AREA)] for i, p in enumerate(self.pagine, 1)],
        )
        scrivi_tabella(
            cartella / "oggetti_visivi",
            ["pagina", "tipo", "titolo"],
            [[p["titolo"], u, n] for p in self.pagine for u, n in p["visual"]],
        )
        scrivi_tabella(
            cartella / "document_properties",
            ["nome", "tipo", "valore", "descrizione", "usata_in", "attributi"],
            [[p["nome"], "standard" if p["standard"] else "utente", p["valore"], p["descrizione"], p["usata_in"], p["attributi"]] for p in self.proprieta_documento],
        )
        scrivi_tabella(
            cartella / "tutte_le_proprieta",
            ["nome", "gruppo", "registro", "standard", "attributi", "valore", "descrizione"],
            [[p["nome"], p["gruppo"], p["classe"], "si" if p["standard"] else "no", p["attributi"], p["valore"], p["descrizione"]] for p in self.proprieta],
        )
        scrivi_tabella(
            cartella / "colonne_calcolate",
            ["tabella", "colonna", "espressione"],
            [[c["tabella"], c["nome"], c["espressione"]] for c in self.colonne_calcolate],
        )
        scrivi_tabella(
            cartella / "data_function",
            ["nome", "linguaggio", "file", "righe_di_codice", "ingressi", "uscite"],
            [[d["nome"], d["linguaggio"], d["file"], d["righe"], ", ".join(d["ingressi"]), ", ".join(d["uscite"])] for d in self.data_function],
        )
        scrivi_tabella(
            cartella / "text_area",
            ["pagina", "titolo", "file", "controlli_spotfire", "immagini", "tag_script_nell_html"],
            [[t["pagina"], t["titolo"], t["file"], t["controlli"], t["immagini"], t["tag_script"]] for t in self.text_area],
        )
        scrivi_tabella(
            cartella / "query_personalizzate",
            ["nome", "file", "righe_sql"],
            [[q["nome"], q["file"], len([r for r in q["sql"].splitlines() if r.strip()])] for q in self.query],
        )
        scrivi_tabella(
            cartella / "tabelle_sorgente",
            ["nome_tabella", "tipo_sorgente", "oggetto_database", "tipo_oggetto_database", "attributi"],
            [[t["nome"], t["natura"], t["oggetto_database"], t["tipo_oggetto_database"], t["dettagli"]] for t in self.tabelle_sorgente],
        )
        scrivi_tabella(
            cartella / "connessioni",
            ["contesto", "chiave", "valore"],
            [[c["contesto"], c["chiave"], c["valore"]] for c in self.connessioni],
        )
        tutti_i_tipi = set(self.tipi_vivi) | set(self.tipi_istantanea)
        scrivi_tabella(
            cartella / "tipi_oggetto",
            ["tipo_completo", "tipo", "nel_documento", "nei_segnalibri"],
            [[t, ultimo_segmento(t), self.tipi_vivi.get(t, 0), self.tipi_istantanea.get(t, 0)] for t in sorted(tutti_i_tipi, key=lambda x: -(self.tipi_vivi.get(x, 0) + self.tipi_istantanea.get(x, 0)))],
        )
        scrivi_tabella(
            cartella / "contenuto_archivio",
            ["voce", "dimensione_byte", "dimensione_compressa_byte"],
            [list(v) for v in self.voci_zip],
        )
        if self.errori:
            scrivi_testo(cartella / "errori.log", "\n".join(self.errori))

    def riepilogo(self):
        text_area = sum(n for t, n in self.visual.items() if t in TIPI_TEXT_AREA)
        totale_visual = sum(self.visual.values())
        return OrderedDict([
            ("dashboard", self.nome),
            ("pagine", len(self.pagine)),
            ("nomi_pagine", ", ".join(p["titolo"] for p in self.pagine)),
            ("oggetti_visivi_totali", totale_visual),
            ("grafici_e_tabelle", totale_visual - text_area),
            ("text_area", text_area),
            ("document_properties", len(self.proprieta_documento)),
            ("document_properties_utente", len(self.proprieta_documento) - len(self.proprieta_documento_standard)),
            ("document_properties_standard", len(self.proprieta_documento_standard)),
            ("script_ironpython", sum(1 for s in self.script if s["linguaggio"] == "IronPython")),
            ("script_python", sum(1 for s in self.script if s["linguaggio"] == "Python")),
            ("script_javascript", sum(1 for s in self.script if s["linguaggio"] == "JavaScript")),
            ("data_function", len(self.data_function)),
            ("query_personalizzate", len(self.query)),
            ("tabelle_o_viste_database", sum(1 for t in self.tabelle_sorgente if t["natura"] == "Tabella o vista del database")),
            ("information_link", sum(1 for t in self.tabelle_sorgente if t["natura"] == "Information Link")),
            ("colonne_calcolate", len(self.colonne_calcolate)),
            ("segnalibri", self.segnalibri),
            ("immagini", len(self.immagini)),
        ])


def main():
    percorsi = [Path(a) for a in sys.argv[1:]] or FILE_DXP
    CARTELLA_OUTPUT.mkdir(parents=True, exist_ok=True)
    if Workbook is None:
        print("openpyxl non installato: output in formato csv. Per ottenere file xlsx eseguire: pip install openpyxl")
    riepiloghi = []
    for percorso in percorsi:
        print(percorso.name)
        if not percorso.exists():
            print("  file non trovato")
            continue
        try:
            analizzatore = AnalizzatoreDxp(percorso)
            analizzatore.esegui()
            analizzatore.esporta(CARTELLA_OUTPUT / pulisci_nome(percorso.stem))
        except zipfile.BadZipFile:
            print("  archivio non leggibile")
            continue
        except PermissionError as e:
            print(f"  accesso negato: {e}")
            continue
        r = analizzatore.riepilogo()
        riepiloghi.append(r)
        print(f"  pagine {r['pagine']}, oggetti visivi {r['oggetti_visivi_totali']}, document properties {r['document_properties']}")
        print(f"  ironpython {r['script_ironpython']}, javascript {r['script_javascript']}, query personalizzate {r['query_personalizzate']}")
        for avviso in analizzatore.avvisi:
            print(f"  {avviso}")
        if analizzatore.errori:
            print(f"  errori: {len(analizzatore.errori)}, vedere errori.log")
    if riepiloghi:
        scrivi_tabella(CARTELLA_OUTPUT / "riepilogo_dashboard", list(riepiloghi[0].keys()), [list(r.values()) for r in riepiloghi])
    print(f"output in {CARTELLA_OUTPUT}")


if __name__ == "__main__":
    main()
