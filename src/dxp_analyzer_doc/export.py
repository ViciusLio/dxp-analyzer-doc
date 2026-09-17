"""Export an :class:`AnalysisResult` to spreadsheet tables and extracted files.

Produces ``.xlsx`` when ``openpyxl`` is installed, otherwise ``.csv``. Column
headers and sheet names are localized (``en`` default, ``it``).
"""

from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import List

from ._utils import clean_name
from .i18n import t
from .model import AnalysisResult
from ._utils import last_segment

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - optional dependency
    Workbook = None

_SCRIPT_DIRS = {"IronPython": ("ironpython", ".py"), "Python": ("python", ".py"), "JavaScript": ("javascript", ".js")}


def _cell_value(value):
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return value
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))
    return text[:32000]


def write_table(path: Path, header, rows) -> Path:
    """Write a table as ``.xlsx`` (if openpyxl) or ``.csv``. Returns the file path."""
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = list(rows)
    if Workbook is None:
        out = path.with_suffix(".csv")
        with open(out, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow(header)
            for row in rows:
                writer.writerow(row)
        return out
    wb = Workbook()
    ws = wb.active
    ws.title = clean_name(path.stem, 31).replace("[", "").replace("]", "")
    ws.append(list(header))
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in rows:
        ws.append([_cell_value(v) for v in row])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, title in enumerate(header, 1):
        widths = [len(str(title))] + [len(str(r[i - 1])) for r in rows[:200] if i - 1 < len(r)]
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, max(widths) + 2))
    out = path.with_suffix(".xlsx")
    wb.save(out)
    return out


def write_text(path: Path, content: str, errors=None) -> bool:
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        with open(path, "w", encoding="utf-8", newline="") as f:
            f.write(content or "")
        return True
    except OSError as e:
        if errors is not None:
            errors.append(f"Impossibile scrivere {path}: {e}")
        return False


def _free_name(name, extension, used) -> str:
    base = clean_name(name)
    candidate = base
    i = 2
    while (candidate + extension).lower() in used:
        candidate = f"{base} ({i})"
        i += 1
    used.add((candidate + extension).lower())
    return candidate + extension


def export_result(result: AnalysisResult, folder, lang: str = "en") -> None:
    """Write every table and extracted code/HTML/SQL file for one dashboard."""
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    yes, no = t("hdr.yes", lang), t("hdr.no", lang)

    # -- extracted code files ---------------------------------------------------
    used = {}
    for s in result.scripts:
        sub, ext = _SCRIPT_DIRS.get(s.language, ("other_scripts", ".txt"))
        name = _free_name(s.name, ext, used.setdefault(sub, set()))
        s.file = f"{sub}/{name}"
        write_text(folder / sub / name, s.code, result.errors)

    used_df = set()
    for d in result.data_functions:
        ext = ".py" if d.language == "Python" else ".R"
        name = _free_name(d.name, ext, used_df)
        d.file = f"data_functions/{name}" if d.code else ""
        if d.code:
            write_text(folder / "data_functions" / name, d.code, result.errors)

    used_html = set()
    for ta in result.text_areas:
        ta.file = ""
        if ta.html.strip():
            name = _free_name(f"{ta.page} - {ta.title}", ".html", used_html)
            ta.file = f"text_area_html/{name}"
            write_text(folder / "text_area_html" / name, ta.html, result.errors)

    used_sql = set()
    for q in result.queries:
        name = _free_name(q.name, ".sql", used_sql)
        q.file = f"query_sql/{name}"
        write_text(folder / "query_sql" / name, q.sql, result.errors)

    # -- tables -----------------------------------------------------------------
    write_table(folder / t("sheet.scripts", lang),
                [t("hdr.name", lang), t("hdr.language", lang), t("hdr.file", lang), t("hdr.code_lines", lang),
                 t("hdr.name_in_file", lang), t("hdr.source_resource", lang)],
                [[s.name, s.language, s.file, s.lines, yes if s.name_in_file else no, s.origin] for s in result.scripts])
    write_table(folder / t("sheet.pages", lang),
                [t("hdr.number", lang), t("hdr.page", lang), t("hdr.visual_objects", lang), t("hdr.text_area", lang), t("hdr.charts", lang)],
                [[i, p.title, len(p.visuals),
                  sum(1 for v in p.visuals if v.type in ("HtmlTextArea", "TextArea")),
                  sum(1 for v in p.visuals if v.type not in ("HtmlTextArea", "TextArea"))] for i, p in enumerate(result.pages, 1)])
    write_table(folder / t("sheet.visual_objects", lang),
                [t("hdr.page", lang), t("hdr.type", lang), t("hdr.title", lang)],
                [[p.title, v.type, v.name] for p in result.pages for v in p.visuals])
    write_table(folder / t("sheet.document_properties", lang),
                [t("hdr.name", lang), t("hdr.type", lang), t("hdr.value", lang), t("hdr.description", lang), t("hdr.used_in", lang), t("hdr.attributes", lang)],
                [[p.name, t(f"kind.{p.kind}", lang), p.value, p.description, p.used_in, p.attributes] for p in result.document_properties])
    write_table(folder / t("sheet.all_properties", lang),
                [t("hdr.name", lang), t("hdr.group", lang), t("hdr.register", lang), t("hdr.standard", lang), t("hdr.attributes", lang), t("hdr.value", lang), t("hdr.description", lang)],
                [[p.name, _group_label(p.group, lang), p.class_name, yes if p.standard else no, p.attributes, p.value, p.description] for p in result.all_properties])
    write_table(folder / t("sheet.calculated_columns", lang),
                [t("hdr.table", lang), t("hdr.column", lang), t("hdr.expression", lang)],
                [[c.table, c.name, c.expression] for c in result.calculated_columns])
    write_table(folder / t("sheet.data_functions", lang),
                [t("hdr.name", lang), t("hdr.language", lang), t("hdr.file", lang), t("hdr.code_lines", lang), t("hdr.inputs", lang), t("hdr.outputs", lang)],
                [[d.name, d.language, d.file, d.lines, ", ".join(d.inputs), ", ".join(d.outputs)] for d in result.data_functions])
    write_table(folder / t("sheet.text_areas", lang),
                [t("hdr.page", lang), t("hdr.title", lang), t("hdr.file", lang), t("hdr.controls_spotfire", lang), t("hdr.images", lang), t("hdr.script_tags", lang)],
                [[ta.page, ta.title, ta.file, ta.controls, ta.images, ta.script_tags] for ta in result.text_areas])
    write_table(folder / t("sheet.custom_queries", lang),
                [t("hdr.name", lang), t("hdr.file", lang), t("hdr.sql_lines", lang)],
                [[q.name, q.file, len([r for r in q.sql.splitlines() if r.strip()])] for q in result.queries])
    write_table(folder / t("sheet.source_tables", lang),
                [t("hdr.table_name", lang), t("hdr.source_type", lang), t("hdr.database_object", lang), t("hdr.database_object_type", lang), t("hdr.attributes", lang)],
                [[s.name, t(f"nature.{s.nature}", lang), s.database_object, s.database_object_type, s.details] for s in result.source_tables])
    write_table(folder / t("sheet.data_tables", lang),
                [t("hdr.table_name", lang)],
                [[name] for name in result.data_tables])
    write_table(folder / t("sheet.connections", lang),
                [t("hdr.context", lang), t("hdr.key", lang), t("hdr.value", lang)],
                [[c.context, c.key, c.value] for c in result.connections])
    all_types = set(result.live_types) | set(result.snapshot_types)
    write_table(folder / t("sheet.object_types", lang),
                [t("hdr.full_type", lang), t("hdr.type", lang), t("hdr.in_document", lang), t("hdr.in_bookmarks", lang)],
                [[ty, last_segment(ty), result.live_types.get(ty, 0), result.snapshot_types.get(ty, 0)]
                 for ty in sorted(all_types, key=lambda x: -(result.live_types.get(x, 0) + result.snapshot_types.get(x, 0)))])
    write_table(folder / t("sheet.archive_content", lang),
                [t("hdr.entry", lang), t("hdr.size_bytes", lang), t("hdr.compressed_bytes", lang)],
                [[z.name, z.size, z.compressed_size] for z in result.zip_entries])
    if result.errors:
        write_text(folder / "errors.log", "\n".join(result.errors))


def write_summary(results: List[AnalysisResult], folder, lang: str = "en") -> None:
    """Write the cross-dashboard summary table."""
    if not results:
        return
    folder = Path(folder)
    metrics = [key for key, _ in results[0].metrics()]
    header = [t(f"metric.{key}", lang) for key in metrics]
    rows = [[value for _, value in r.metrics()] for r in results]
    write_table(folder / t("sheet.summary", lang), header, rows)


def _group_label(group: str, lang: str) -> str:
    if group in ("document", "column", "table"):
        return t(f"group.{group}", lang)
    return group
